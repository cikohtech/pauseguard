"""Fill the official AI Builder Cup template with PauseGuard's content (phase 11; `make deck`).

The template (`reference/Submission_Template_AI_Builder_Cup.pptx`) keeps its slides, order,
titles, header/footer bars and fonts; this script only adds text boxes, pictures and one table to
them, and fills in the text boxes the template already has (team details, opportunities, links).
The text comes from `docs/SUBMISSION_DECK.md` (bullets and tables under each "## N." heading), the
pictures from `docs/img/` and `eval/reports/charts/`. Missing pictures become a labelled grey box
and a warning, so a draft can be built before `make screens`; `--strict` turns that into an error.

    make deck [VIDEO_URL=https://youtu.be/...] [STRICT=1]
"""

from __future__ import annotations

import argparse
import copy
import io
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "reference" / "Submission_Template_AI_Builder_Cup.pptx"
CONTENT = ROOT / "docs" / "SUBMISSION_DECK.md"
OUT = ROOT / "submission" / "PauseGuard_Submission.pptx"
SCREENS = ROOT / "docs" / "img" / "screens"

FONT = "Google Sans Flex"  # the template's font family
INK = RGBColor(0x20, 0x21, 0x24)
MUTED = RGBColor(0x5F, 0x63, 0x68)
PANEL = RGBColor(0xF1, 0xF3, 0xF4)
LINE = RGBColor(0xDA, 0xDC, 0xE0)
HEADER = RGBColor(0x20, 0x21, 0x24)
GREEN = RGBColor(0x18, 0x80, 0x38)
LINK = RGBColor(0x1A, 0x73, 0xE8)

# Pictures are downscaled to this many pixels per inch of their size on the slide (sharp on a
# projector, and it keeps the deck under the repo's 2 MB file limit).
PPI = 200

# The content area under the template's header bar and slide title, above the footer bar (inches).
LEFT, TOP, WIDTH, BOTTOM = 0.32, 1.28, 9.36, 5.42

missing: list[str] = []


# --------------------------------------------------------------------------------------------
# docs/SUBMISSION_DECK.md
# --------------------------------------------------------------------------------------------


@dataclass
class Section:
    title: str
    items: list[tuple[int, str]] = field(default_factory=list)  # (level, text)
    table: list[list[str]] = field(default_factory=list)


def parse_content(path: Path) -> dict[int, Section]:
    sections: dict[int, Section] = {}
    current: Section | None = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"^## (\d+)\. (.+)$", raw)
        if heading:
            current = sections.setdefault(int(heading.group(1)), Section(heading.group(2).strip()))
            continue
        if raw.startswith("## "):
            current = None
            continue
        if current is None or not raw.strip() or raw.lstrip().startswith(">"):
            continue
        bullet = re.match(r"^( *)- (.+)$", raw)
        if bullet:
            current.items.append((len(bullet.group(1)) // 2, bullet.group(2).strip()))
        elif raw.startswith("|"):
            cells = [c.strip() for c in raw.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                current.table.append(cells)
    return sections


def runs_of(text: str) -> list[tuple[str, bool]]:
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


# --------------------------------------------------------------------------------------------
# Drawing helpers
# --------------------------------------------------------------------------------------------


def style_run(run, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def set_bullet(paragraph, char: str | None, level_indent: float, hang: float = 0.18) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buSzPct"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    if char is None:
        ppr.set("marL", str(Inches(level_indent)))
        ppr.set("indent", "0")
        ppr.append(ppr.makeelement(qn("a:buNone"), {}))
        return
    ppr.set("marL", str(Inches(level_indent + hang)))
    ppr.set("indent", str(-Inches(hang)))
    ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": FONT}))
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": char}))


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paragraphs: list[tuple[int, str]],
    size: float = 12,
    sub_size: float | None = None,
    bullets: bool = True,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    color: RGBColor = INK,
    space_after: float = 4,
    first_bold: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, (level, text) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        s = size if level == 0 else (sub_size or size - 1)
        for chunk, bold in runs_of(text):
            run = p.add_run()
            run.text = chunk
            style_run(run, s, bold or (first_bold and i == 0), color)
        if bullets:
            set_bullet(p, "●" if level == 0 else "–", 0.0 if level == 0 else 0.22)
        else:
            set_bullet(p, None, 0.0)
    return box


def add_panel(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor = PANEL,
    line: RGBColor | None = LINE,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def scaled_picture(path: Path, width_in: float) -> io.BytesIO | str:
    """The picture at PPI for its width on the slide; app screenshots as JPEG, diagrams as PNG."""
    from PIL import Image

    screenshot = path.parent == SCREENS
    with Image.open(path) as im:
        target = round(width_in * PPI)
        if im.width <= target * 1.1 and not screenshot:
            return str(path)
        size = (target, round(im.height * target / im.width)) if im.width > target else im.size
        small = im.convert("RGB").resize(size, Image.LANCZOS)
    buf = io.BytesIO()
    if screenshot:
        small.save(buf, format="JPEG", quality=88, subsampling=0, optimize=True)
    else:
        small.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    return buf


def image_size(path: Path) -> tuple[int, int]:
    from PIL import Image

    with Image.open(path) as im:
        return im.size


def add_image(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    align: str = "center",
    border: bool = False,
):
    """Fit `path` inside the box, keeping its aspect ratio; returns (left, top, w, h) in inches."""
    if not path.exists():
        missing.append(str(path.relative_to(ROOT)))
        ph = add_panel(slide, x, y, w, h, fill=RGBColor(0xE8, 0xEA, 0xED), line=None)
        tf = ph.text_frame
        tf.text = f"missing: {path.relative_to(ROOT)}"
        style_run(tf.paragraphs[0].runs[0], 9, color=MUTED)
        return x, y, w, h
    iw, ih = image_size(path)
    scale = min(w / iw, h / ih)
    pw, ph_ = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph_) / 2 if align == "center" else y
    pic = slide.shapes.add_picture(
        scaled_picture(path, pw), Inches(px), Inches(py), Inches(pw), Inches(ph_)
    )
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return px, py, pw, ph_


def title_box(slide):
    """The template's slide-title text box (the first text box with text below the header bar)."""
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() and Emu(shape.top).inches > 0.6:
            return shape
    raise ValueError("no title box on this slide")


def add_hyperlink_run(paragraph, text: str, url: str, size: float, like) -> None:
    run = paragraph.add_run()
    run.text = text
    if like is not None and like.font.name:
        run.font.name = like.font.name
    run.font.size = Pt(size)
    run.font.color.rgb = LINK
    if url.startswith("http"):
        run.hyperlink.address = url


# --------------------------------------------------------------------------------------------
# Slides (template order: 1 team, 2 organisers' note, 3 idea, 4 opportunities, 5 features,
# 6 process flow, 7 wireframes, 8 architecture, 9 technologies, 10 cost, 11 snapshots,
# 12 benchmark, 13 future, 14 links, 15 spare -> the four mandatory technologies, 16 thank you)
# --------------------------------------------------------------------------------------------


def fill_team(slide, sec: Section) -> None:
    values = {}
    for _, text in sec.items:
        label, _, value = text.partition(":")
        values[label.strip().lower()] = value.strip()
    box = next(s for s in slide.shapes if s.has_text_frame and "Team Details" in s.text_frame.text)
    paragraphs = box.text_frame.paragraphs
    for p in paragraphs:
        label = p.text.strip().rstrip(":").lower()
        if label in ("team name", "team leader name") and values.get(label):
            run = p.add_run()
            run.text = values[label]
            style_run(run, 18, bold=True)
        if label == "problem statement" and values.get(label):
            new = copy.deepcopy(p._p)
            p._p.addnext(new)
            from pptx.text.text import _Paragraph

            q = _Paragraph(new, p._parent)
            for r in list(q.runs):
                r._r.getparent().remove(r._r)
            set_bullet(q, None, 0.75)
            style_run(q.add_run(), 11, color=INK)
            q.runs[0].text = values[label]
    # The template's trailing empty paragraph only pushes the text down.
    last = box.text_frame.paragraphs[-1]
    if not last.text.strip() and len(box.text_frame.paragraphs) > 3:
        last._p.getparent().remove(last._p)


def fill_idea(slide, sec: Section) -> None:
    add_text(
        slide, LEFT, TOP + 0.05, 6.15, BOTTOM - TOP - 0.05, sec.items, size=12.5, space_after=7
    )
    add_image(slide, SCREENS / "04-red.png", 6.65, TOP - 0.05, 3.0, BOTTOM - TOP + 0.05)


def fill_opportunities(slide, sec: Section) -> None:
    box = title_box(slide)
    tf = box.text_frame
    answers: dict[str, list[str]] = {}
    question = None
    for level, text in sec.items:
        if level == 0:
            question = text
            answers[question] = []
        elif question:
            answers[question].append(text)
    for p in list(tf.paragraphs):
        q = p.text.strip()
        if q not in answers:
            continue
        anchor = p._p
        for text in reversed(answers[q]):
            new = copy.deepcopy(p._p)
            anchor.addnext(new)
            from pptx.text.text import _Paragraph

            para = _Paragraph(new, p._parent)
            for r in list(para.runs):
                r._r.getparent().remove(r._r)
            for chunk, bold in runs_of(text):
                run = para.add_run()
                run.text = chunk
                style_run(run, 12.5, bold, INK)
            set_bullet(para, None, 0.5)
            para.space_after = Pt(9)
        for r in p.runs:
            style_run(r, 13, True, INK)
    box.height = Inches(BOTTOM - Emu(box.top).inches)


def fill_features(slide, sec: Section) -> None:
    groups: list[list[tuple[int, str]]] = []
    for level, text in sec.items:
        if level == 0:
            groups.append([(0, text)])
        else:
            groups[-1].append((1, text))
    col_w = (WIDTH - 0.25) / 2
    for i, group in enumerate(groups[:2]):
        x = LEFT + i * (col_w + 0.25)
        add_panel(slide, x, TOP, col_w, BOTTOM - TOP)
        header = [(0, group[0][1])]
        add_text(
            slide,
            x + 0.15,
            TOP + 0.1,
            col_w - 0.3,
            0.4,
            header,
            size=13,
            bullets=False,
            first_bold=True,
        )
        items = [(0, t) for _, t in group[1:]]
        add_text(
            slide,
            x + 0.15,
            TOP + 0.5,
            col_w - 0.3,
            BOTTOM - TOP - 0.6,
            items,
            size=11.5,
            space_after=5,
        )


def fill_image_slide(slide, path: Path) -> None:
    add_image(slide, path, LEFT, TOP, WIDTH, BOTTOM - TOP)


def fill_wireframes(slide, sec: Section) -> None:
    shots = [
        "01-persona-picker.png",
        "02-checkin.png",
        "03-amber.png",
        "04-red.png",
        "04-red-zh.png",
    ]
    captions = [t for _, t in sec.items]
    n = len(shots)
    gap = 0.17
    w = (WIDTH - gap * (n - 1)) / n
    for i, (shot, caption) in enumerate(zip(shots, captions, strict=False)):
        x = LEFT + i * (w + gap)
        _, py, _, ph = add_image(slide, SCREENS / shot, x, TOP, w, 3.55, border=True)
        add_text(
            slide,
            x - 0.05,
            TOP + 3.62,
            w + 0.1,
            0.5,
            [(0, caption)],
            size=9,
            bullets=False,
            align=PP_ALIGN.CENTER,
            color=MUTED,
        )


def fill_architecture(slide, sec: Section) -> None:
    add_image(slide, ROOT / "docs" / "img" / "architecture.png", LEFT, TOP - 0.08, WIDTH, 3.12)
    n = len(sec.items)
    gap = 0.15
    w = (WIDTH - gap * (n - 1)) / n
    y = TOP + 3.1
    for i, (_, text) in enumerate(sec.items):
        x = LEFT + i * (w + gap)
        add_panel(slide, x, y, w, BOTTOM - y)
        add_text(
            slide,
            x + 0.08,
            y + 0.04,
            w - 0.16,
            BOTTOM - y - 0.08,
            [(0, text)],
            size=9,
            bullets=False,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def fill_bullets(
    slide, sec: Section, size: float = 12, sub: float | None = None, space: float = 6
) -> None:
    add_text(
        slide, LEFT, TOP, WIDTH, BOTTOM - TOP, sec.items, size=size, sub_size=sub, space_after=space
    )


def fill_snapshots(slide, sec: Section) -> None:
    captions = [t for _, t in sec.items]
    right_w = 2.95
    big_w = WIDTH - right_w - 0.2
    # Left: the console, as large as the slide allows; right: two smaller shots stacked.
    _, py, pw, ph = add_image(
        slide, SCREENS / "05-console-evidence.png", LEFT, TOP, big_w, 3.62, align="top", border=True
    )
    add_text(
        slide,
        LEFT,
        py + ph + 0.05,
        pw,
        0.5,
        [(0, captions[0])],
        size=9,
        bullets=False,
        align=PP_ALIGN.CENTER,
        color=MUTED,
    )
    x = LEFT + big_w + 0.2
    y = TOP
    for path, caption in (
        (SCREENS / "06-dashboard.png", captions[1]),
        (SCREENS / "00-landing.png", captions[2]),
    ):
        _, py, pw, ph = add_image(slide, path, x, y, right_w, 1.84, align="top", border=True)
        add_text(
            slide,
            x - 0.05,
            py + ph + 0.03,
            right_w + 0.1,
            0.32,
            [(0, caption)],
            size=8.5,
            bullets=False,
            align=PP_ALIGN.CENTER,
            color=MUTED,
        )
        y = py + ph + 0.36


def fill_benchmark(slide, sec: Section) -> None:
    intro, rest = sec.items[:1], sec.items[1:]
    add_text(slide, LEFT, TOP - 0.05, WIDTH, 0.5, intro, size=10, bullets=False, color=MUTED)
    rows, cols = len(sec.table), len(sec.table[0])
    tw, rh = 6.0, 0.27
    shape = slide.shapes.add_table(
        rows, cols, Inches(LEFT), Inches(TOP + 0.5), Inches(tw), Inches(rh * rows)
    )
    table = shape.table
    tbl_pr = table._tbl.tblPr
    for el in tbl_pr.findall(qn("a:tableStyleId")):
        tbl_pr.remove(el)
    widths = [1.55] + [(tw - 1.55) / (cols - 1)] * (cols - 1)
    for c, cw in enumerate(widths):
        table.columns[c].width = Inches(cw)
    for r, row in enumerate(sec.table):
        table.rows[r].height = Inches(rh)
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_ours = "PauseGuard" in row[0]
            cell.fill.fore_color.rgb = (
                HEADER
                if r == 0
                else (RGBColor(0xE6, 0xF4, 0xEA) if is_ours else RGBColor(0xFF, 0xFF, 0xFF))
            )
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            for chunk, bold in runs_of(value):
                run = p.add_run()
                run.text = chunk
                style_run(
                    run,
                    8.5 if r == 0 else 9,
                    bold or r == 0,
                    RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else INK,
                )
    add_image(
        slide,
        ROOT / "eval" / "reports" / "charts" / "benchmark_recall_fpr.png",
        LEFT + tw + 0.15,
        TOP + 0.5,
        WIDTH - tw - 0.15,
        BOTTOM - TOP - 0.5,
        align="top",
    )
    y = TOP + 0.5 + rh * rows + 0.15
    add_text(slide, LEFT, y, tw, BOTTOM - y, rest, size=9.5, space_after=3)


def fill_links(slide, sec: Section, video_url: str | None) -> None:
    box = next(s for s in slide.shapes if s.has_text_frame and "Provide links" in s.text_frame.text)
    values = {}
    for _, text in sec.items:
        label, _, value = text.partition(":")
        values[label.strip()] = value.strip()
    for p in box.text_frame.paragraphs:
        label = p.text.strip()
        if label not in values:
            continue
        value = values[label]
        like = p.runs[0] if p.runs else None
        if value == "VIDEO_URL" and not video_url:
            sep = p.add_run()
            sep.text = ": (the YouTube link is added before submitting)"
            style_run(sep, 12, color=MUTED)
            p.space_after = Pt(10)
            continue
        if value == "VIDEO_URL":
            value = video_url
        sep = p.add_run()
        sep.text = ": "
        if like is not None:
            sep.font.name = like.font.name
        sep.font.size = Pt(14)
        url = value.split(" ")[0]
        add_hyperlink_run(p, url, url, 14, like)
        extra = value[len(url) :].strip()
        if extra:
            r = p.add_run()
            r.text = " " + extra
            style_run(r, 11, color=MUTED)
        p.space_after = Pt(10)
    box.height = Inches(3.6)


def fill_spare(slide, sec: Section, template_title) -> None:
    """Slide 15 is an empty content slide in the template: give it a title like the others."""
    new = copy.deepcopy(template_title._element)
    slide.shapes._spTree.append(new)
    shape = slide.shapes[-1]
    runs = shape.text_frame.paragraphs[0].runs
    runs[0].text = sec.title
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    gap = 0.2
    w = (WIDTH - gap) / 2
    h = (BOTTOM - TOP - gap) / 2
    for i, (_, text) in enumerate(sec.items[:4]):
        x = LEFT + (i % 2) * (w + gap)
        y = TOP + (i // 2) * (h + gap)
        add_panel(slide, x, y, w, h)
        add_text(
            slide,
            x + 0.15,
            y + 0.08,
            w - 0.3,
            h - 0.16,
            [(0, text)],
            size=12.5,
            bullets=False,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def build(video_url: str | None, strict: bool) -> Path:
    content = parse_content(CONTENT)
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)
    if len(slides) != 16:
        raise SystemExit(f"expected the 16-slide official template, got {len(slides)} slides")

    fill_team(slides[0], content[1])
    # slides[1] is the organisers' note: kept as is (the team's choice).
    fill_idea(slides[2], content[3])
    fill_opportunities(slides[3], content[4])
    fill_features(slides[4], content[5])
    fill_image_slide(slides[5], ROOT / "docs" / "img" / "process-flow.png")
    fill_wireframes(slides[6], content[7])
    fill_architecture(slides[7], content[8])
    fill_bullets(slides[8], content[9], size=14, sub=12.5, space=8)
    fill_bullets(slides[9], content[10], size=13, sub=12, space=8)
    fill_snapshots(slides[10], content[11])
    fill_benchmark(slides[11], content[12])
    fill_bullets(slides[12], content[13], size=14, space=11)
    fill_links(slides[13], content[14], video_url)
    fill_spare(slides[14], content[15], title_box(slides[2]))

    if missing:
        print("⚠ missing pictures (grey boxes in the deck):", *missing, sep="\n  ", file=sys.stderr)
        if strict:
            raise SystemExit("--strict: run `make screens` first")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--video-url", default=None, help="the YouTube link for the links slide")
    ap.add_argument("--strict", action="store_true", help="fail if a picture is missing")
    args = ap.parse_args()
    out = build(args.video_url, args.strict)
    print(f"✔ {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
